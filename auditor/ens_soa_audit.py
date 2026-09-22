#!/usr/bin/env python3
"""ens-soa-audit · Auditor de Declaraciones de Aplicabilidad del ENS (RD 311/2022).

Lee una SoA en Excel (plantilla del caso TechServ o la exportada por ENS Compliance Studio),
recalcula de forma independiente el nivel exigido de cada medida a partir de la hoja de
categorización y comprueba la coherencia interna del documento con las mismas reglas
(y los mismos identificadores) que el auditor de la aplicación web.

    python3 ens_soa_audit.py SoA.xlsx                     # resumen en consola
    python3 ens_soa_audit.py SoA.xlsx --md informe.md     # informe Markdown
    python3 ens_soa_audit.py SoA.xlsx --json out.json     # resultado estructurado
    python3 ens_soa_audit.py SoA.xlsx --fail-on mayor     # código de salida 2 si hay NC mayores (CI)

Solo depende de openpyxl. Herramienta de preauditoría: no sustituye a la auditoría formal (art. 31).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Falta openpyxl: pip install openpyxl")

__version__ = "1.0.0"
HERE = Path(__file__).resolve().parent
ANEXO = json.loads((HERE / "anexo_ii.json").read_text(encoding="utf-8"))["medidas"]

DIMS = ["D", "I", "C", "A", "T"]
DIM_LABEL = {"D": "Disponibilidad", "I": "Integridad", "C": "Confidencialidad", "A": "Autenticidad", "T": "Trazabilidad"}
LV = {"BAJO": 1, "MEDIO": 2, "ALTO": 3}
CAT_TO_LEVEL = {"ALTA": "ALTO", "MEDIA": "MEDIO", "BÁSICA": "BAJO"}
MAYOR, MENOR, OBS = "NC mayor", "NC menor", "Observación"
SEV_ORDER = {MAYOR: 0, MENOR: 1, OBS: 2}
ERROR_VALUES = {"#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}
CODE_RE = re.compile(r"\b(?:org\.\d+|(?:op|mp)\.[a-z]+\.\d+)\b")

# Cabeceras canónicas de la hoja SoA (plantilla TechServ / exportación de la app) → clave interna
SOA_COLUMNS = {
    "Marco / familia": "familia", "Código": "codigo", "Medida de seguridad": "nombre", "Dimensiones afectadas (Anexo II)": "dims",
    "Nivel BAJO": "bajo", "Nivel MEDIO": "medio", "Nivel ALTO": "alto", "Nivel exigido": "nivel_exigido",
    "Regla aplicada": "regla", "Exigencia aplicable (Anexo II)": "exigencia", "Refuerzos exigidos (Rn)": "refuerzos_exigidos",
    "Refuerzos / alternativa elegida": "refuerzos_elegidos", "¿Aplica?": "aplica", "Justificación de aplicabilidad / exclusión": "justificacion",
    "Medidas ORGANIZATIVAS implantadas": "org", "Medidas TÉCNICAS implantadas": "tec", "Medida compensatoria (ref.)": "mc_ref",
    "Estado de implantación": "estado", "% implantación": "pct", "Evidencias / documentos": "evidencias", "Responsable": "responsable",
    "Control ISO/IEC 27001:2022 equivalente (CCN-STIC 825)": "iso27001", "Guías CCN-STIC / verificación": "guias", "Observaciones / PTR": "observaciones",
}
REQUIRED = ["codigo", "aplica", "estado", "evidencias", "responsable"]
TOLERATED_SUFFIX = {"nivel_exigido"}  # cabeceras que la plantilla personaliza con el nombre de la entidad


@dataclass
class Finding:
    id: str
    sev: str
    ambito: str
    titulo: str
    detalle: str
    recomendacion: str
    ref: str


@dataclass
class Result:
    fichero: str
    fecha: str
    categoria: str | None = None
    niveles: dict = field(default_factory=dict)
    medidas: int = 0
    aplicables: int = 0
    grado: float | None = None
    hallazgos: list = field(default_factory=list)

    def count(self, sev: str) -> int:
        return sum(1 for f in self.hallazgos if f.sev == sev)


# ---------------------------------------------------------------- utilidades
def norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9%?]+", " ", s).strip()


def blank(v) -> bool:
    return v is None or str(v).strip() in ("", "—", "-")


ALIAS = {
    "B": "BAJO", "BAJ": "BAJO", "BAJA": "BAJO", "BASICA": "BAJO", "BÁSICA": "BAJO",
    "M": "MEDIO", "MED": "MEDIO", "MEDIA": "MEDIO",
    "ALT": "ALTO", "ALTA": "ALTO",
}


def ens_level(v):
    """None = vacío · 'BAJO'|'MEDIO'|'ALTO' = válido (también B/M/ALT) · False = valor no válido."""
    if blank(v):
        return None
    s = str(v).strip().upper().rstrip(".")
    if s in LV:
        return s
    return ALIAS.get(s, False)


def max_level(a, b):
    if not a:
        return b
    if not b:
        return a
    return a if LV[a] >= LV[b] else b


def parse_exigencia(txt):
    s = str(txt or "").strip()
    if not s or re.fullmatch(r"n\.?a\.?", s, re.I):
        return {"aplica": False, "obligatorios": [], "grupos": []}
    grupos = [[x.strip() for x in re.split(r"\s+o\s+", g) if x.strip()] for g in re.findall(r"\[([^\]]+)\]", s)]
    resto = re.sub(r"\[[^\]]+\]", "", s)
    return {"aplica": True, "obligatorios": re.findall(r"R\d+", resto), "grupos": grupos}


def pct_of(row) -> float:
    p = row.get("pct")
    if isinstance(p, (int, float)):
        return float(p)
    if not blank(p):
        try:
            v = float(str(p).replace("%", "").replace(",", "."))
            return v / 100 if v > 1 else v
        except ValueError:
            pass
    return 1.0 if re.fullmatch(r"implantada", str(row.get("estado") or ""), re.I) else 0.0


def aplica_declarada(row) -> bool:
    return bool(re.match(r"^s[ií]", str(row.get("aplica") or "").strip(), re.I))


# ---------------------------------------------------------------- lectura del libro
def find_sheet(wb, *names, header_has=()):
    for n in names:
        for ws in wb.worksheets:
            if norm(ws.title) == norm(n):
                return ws
    if header_has:
        for ws in wb.worksheets:
            for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
                vals = {norm(c) for c in r if c is not None}
                if all(norm(h) in vals for h in header_has):
                    return ws
    return None


def header_row(ws, must=("codigo",), search=8):
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=search, values_only=True), start=1):
        vals = [norm(c) for c in r]
        if all(m in vals for m in must):
            return i, [c for c in r]
    return None, None


def load_workbook(path: Path):
    wb_v = openpyxl.load_workbook(path, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    return wb_v, wb_f


def read_categorizacion(ws):
    hr, hdr = header_row(ws, ("id", "d", "i", "c", "a", "t"))
    if not hr:
        return []
    idx = {norm(h): j for j, h in enumerate(hdr) if h is not None}
    out = []
    for r in ws.iter_rows(min_row=hr + 1, values_only=True):
        rid = r[idx["id"]] if idx.get("id") is not None else None
        if blank(rid) or not re.match(r"^[A-Za-z]+-\d+", str(rid)):
            continue
        name_col = next((j for k, j in idx.items() if k.startswith("activo esencial")), None)
        out.append({"id": str(rid).strip(), "nombre": r[name_col] if name_col is not None else "", **{d: r[idx[d.lower()]] for d in DIMS}})
    return out


def read_soa(ws_v, ws_f, res: Result):
    hr, hdr = header_row(ws_v, ("codigo", "¿aplica?".replace("¿", "")))
    if not hr:
        hr, hdr = header_row(ws_v, ("codigo",))
    if not hr:
        raise SystemExit(f"No se encuentra la fila de cabecera con 'Código' en la hoja {ws_v.title!r}")
    canon = {norm(k): v for k, v in SOA_COLUMNS.items()}
    cols = {}
    for j, h in enumerate(hdr):
        if h is None:
            continue
        n = norm(h)
        key = canon.get(n)
        if key is None:  # cabecera alterada: se acepta si empieza por la canónica
            for cn, ck in canon.items():
                if n.startswith(cn) and ck not in cols:
                    key = ck
                    if ck in TOLERATED_SUFFIX:  # p. ej. «Nivel exigido a TechServ»: variante legítima
                        break
                    res.hallazgos.append(Finding("STR-01", OBS, f"Cabecera col. {openpyxl.utils.get_column_letter(j + 1)}", "Cabecera de columna alterada",
                                                 f"La cabecera '{str(h)[:90]}' contiene texto ajeno a la plantilla (esperado: '{next(k for k, v in SOA_COLUMNS.items() if v == ck)}').",
                                                 "Restaurar la cabecera; un texto pegado en la cabecera suele indicar una edición accidental.", "Control documental"))
                    break
        if key and key not in cols:
            cols[key] = j
    missing = [k for k in REQUIRED if k not in cols]
    if missing:
        raise SystemExit(f"Faltan columnas obligatorias en la SoA: {', '.join(missing)}")
    rows = []
    for (rv, rf) in zip(ws_v.iter_rows(min_row=hr + 1), ws_f.iter_rows(min_row=hr + 1)):
        code = rv[cols["codigo"]].value
        if blank(code) or not CODE_RE.fullmatch(str(code).strip()):
            continue
        row = {k: rv[j].value for k, j in cols.items()}
        row["codigo"] = str(code).strip()
        row["_fila"] = rv[0].row
        for c in rv:
            if isinstance(c.value, str) and c.value.strip() in ERROR_VALUES:
                res.hallazgos.append(Finding("STR-03", MENOR, f"{row['codigo']} ({c.coordinate})", "Error de fórmula en la celda",
                                             f"La celda {c.coordinate} contiene {c.value}.", "Reparar la fórmula o la referencia rota.", "Control documental"))
        rows.append(row)
    return rows


def read_refuerzos(ws):
    if ws is None:
        return None
    hr, hdr = header_row(ws, ("codigo medida", "refuerzo"))
    if not hr:
        return None
    idx = {norm(h): j for j, h in enumerate(hdr) if h is not None}
    ex = next((j for k, j in idx.items() if k.startswith("exigible")), None)
    imp = next((j for k, j in idx.items() if k.startswith("implementacion")), None)
    out = []
    for r in ws.iter_rows(min_row=hr + 1, values_only=True):
        if blank(r[idx["codigo medida"]]):
            continue
        out.append({"codigo": str(r[idx["codigo medida"]]).strip(), "refuerzo": str(r[idx["refuerzo"]]).strip(),
                    "exigible": r[ex] if ex is not None else "", "implementacion": r[imp] if imp is not None else ""})
    return out


def read_mc(ws):
    if ws is None:
        return []
    hr, hdr = header_row(ws, ("id",))
    if not hr:
        return []
    keys = {"id": "id", "medida / refuerzo ens sustituido": "medida", "riesgo identificado": "riesgo", "validacion de la medida compensatoria": "validacion",
            "validacion": "validacion", "mantenimiento / revision": "mantenimiento", "aprobacion": "aprobacion"}
    idx = {}
    for j, h in enumerate(hdr):
        n = norm(h).replace(" ", " ")
        for k, v in keys.items():
            if n == norm(k) or (v == "riesgo" and n.startswith("riesgo identificado")):
                idx[v] = j
    out = []
    for r in ws.iter_rows(min_row=hr + 1, values_only=True):
        if blank(r[idx["id"]]) or not re.match(r"^MC-\d+", str(r[idx["id"]])):
            continue
        out.append({k: r[j] for k, j in idx.items()})
    return out


def read_portada(ws):
    if ws is None:
        return {}
    out = {}
    for r in ws.iter_rows(values_only=True):
        vals = [c for c in r if c is not None]
        if len(vals) == 2:
            out[str(vals[0]).strip()] = vals[1]
    return out


# ---------------------------------------------------------------- auditoría
def audit(path: Path, hoy: dt.date | None = None) -> Result:
    hoy = hoy or dt.date.today()
    res = Result(fichero=path.name, fecha=hoy.isoformat())
    add = lambda *a: res.hallazgos.append(Finding(*a))
    wb_v, wb_f = load_workbook(path)

    ws_cat = find_sheet(wb_v, "Categorización", header_has=("ID", "D", "I", "C", "A", "T"))
    ws_soa_v = find_sheet(wb_v, "SoA ENS", header_has=("Código", "¿Aplica?"))
    if ws_soa_v is None:
        raise SystemExit("El libro no contiene una hoja de SoA reconocible (cabeceras 'Código' y '¿Aplica?').")
    ws_soa_f = wb_f[ws_soa_v.title]
    ws_ref = find_sheet(wb_v, "Refuerzos")
    ws_mc = find_sheet(wb_v, "Medidas compensatorias")
    ws_port = find_sheet(wb_v, "Portada")

    # 1. Categorización (Anexo I)
    cat = read_categorizacion(ws_cat) if ws_cat is not None else []
    niveles = {d: None for d in DIMS}
    invalidos = []
    for a in cat:
        for d in DIMS:
            lv = ens_level(a[d])
            if lv is False:
                invalidos.append((a, d))
            elif lv:
                niveles[d] = max_level(niveles[d], lv)
    vals = [niveles[d] for d in DIMS if niveles[d]]
    categoria = "ALTA" if "ALTO" in vals else "MEDIA" if "MEDIO" in vals else "BÁSICA"
    res.categoria, res.niveles = (categoria if cat else None), niveles
    for a, d in invalidos:
        otros = [b["id"] for b in cat if b is not a and ens_level(b[d]) == "ALTO"]
        impacto = (f"No altera el resultado: {DIM_LABEL[d]} ya es ALTO por {', '.join(otros)}." if otros else
                   f"Puede alterar el resultado: si el valor correcto fuera ALTO, {DIM_LABEL[d]} pasaría de {niveles[d] or 'sin valorar'} a ALTO.")
        add("CAT-01", MENOR, f"{a['id']} · {d}", "Valor de categorización no válido",
            f"El activo esencial {a['id']} ({a['nombre']}) tiene '{a[d]}' en la dimensión {DIM_LABEL[d]}. Solo se admiten BAJO, MEDIO o ALTO; la celda se ignora en el cálculo. {impacto}",
            "Corregir la valoración con criterio del Responsable de la Información (CCN-STIC 803) y volver a aprobar la categorización en el CSI.", "RD 311/2022, Anexo I")
    for a in cat:
        if all(blank(a[d]) for d in DIMS):
            add("CAT-02", MAYOR, a["id"], "Activo esencial sin valorar", f"{a['id']} no tiene valoración en ninguna dimensión.", "Valorar las cinco dimensiones del activo esencial.", "RD 311/2022, art. 40")
    if not cat:
        add("CAT-02", MAYOR, "Categorización", "No se encuentra la categorización", "El libro no incluye la valoración de activos esenciales; no se puede verificar el nivel exigido.", "Incluir la hoja de categorización (Anexo I).", "RD 311/2022, art. 40")

    # 2. SoA
    rows = read_soa(ws_soa_v, ws_soa_f, res)
    seen = {}
    for r in rows:
        if r["codigo"] in seen:
            add("STR-02", MAYOR, r["codigo"], "Medida duplicada", f"{r['codigo']} aparece en las filas {seen[r['codigo']]} y {r['_fila']}.", "Dejar una única declaración por medida.", "RD 311/2022, art. 28.2")
        seen.setdefault(r["codigo"], r["_fila"])
    by_code = {r["codigo"]: r for r in rows}
    refuerzos = read_refuerzos(ws_ref)
    mcs = read_mc(ws_mc)
    res.medidas = len(by_code)
    pcts = []
    for code in sorted(ANEXO, key=code_key):
        ax = ANEXO[code]
        if ax["dims"] == "Categoría":
            nivel, regla = CAT_TO_LEVEL[categoria], f"Categoría del sistema = {categoria}"
        else:
            ds = [c for c in ax["dims"] if c in DIMS]
            nivel = None
            for d in ds:
                nivel = max_level(nivel, niveles[d] or "BAJO")
            nivel = nivel or "BAJO"
            regla = "máx(" + ", ".join(f"{d}={niveles[d] or 'BAJO'}" for d in ds) + f") = {nivel}"
        exig = ax[nivel.lower()]
        pe = parse_exigencia(exig)
        r = by_code.get(code)
        if r is None:
            add("SOA-03", MAYOR, code, "Medida del Anexo II ausente en la SoA", f"{code} {ax['nombre']} no figura en la Declaración de Aplicabilidad.", "La SoA debe pronunciarse sobre las 73 medidas del Anexo II.", "RD 311/2022, art. 28.2")
            continue
        if pe["aplica"]:
            res.aplicables += 1
            pcts.append(pct_of(r) if aplica_declarada(r) else 0.0)
        decl_nivel = r.get("nivel_exigido")
        if not blank(decl_nivel) and str(decl_nivel).strip().upper() in LV and str(decl_nivel).strip().upper() != nivel:
            add("SOA-09", MAYOR, code, "Nivel exigido declarado distinto del calculado",
                f"La SoA indica '{decl_nivel}' para {code}, pero con la categorización del propio libro corresponde {nivel} ({regla}).",
                "Recalcular el nivel exigido tras cualquier cambio de categorización.", "RD 311/2022, Anexo II, apdo. 2")
        dec = aplica_declarada(r)
        if pe["aplica"] and not dec:
            add("SOA-01", MAYOR, code, "Exclusión indebida", f"{code} se declara NO aplicable, pero con nivel {nivel} ({regla}) el Anexo II exige '{exig}'.",
                "Declarar la medida como aplicable e implantarla, o justificar una medida compensatoria (art. 28.3).", "RD 311/2022, art. 28")
        if not pe["aplica"] and dec:
            add("SOA-02", OBS, code, "Medida declarada aplicable sin exigencia normativa", f"{code} no se exige en nivel {nivel} ('n.a.'), pero la SoA la declara aplicable.",
                "Es admisible si responde al análisis de riesgos; dejarlo justificado.", "RD 311/2022, art. 28.1")
        if not pe["aplica"] and not dec and blank(r.get("justificacion")):
            add("SOA-08", MENOR, code, "Exclusión sin justificar", f"{code} se excluye sin justificación documentada.", "Documentar el motivo de la exclusión.", "RD 311/2022, art. 28.2")
        if not dec:
            continue
        if blank(r.get("evidencias")):
            add("SOA-04", MENOR, code, "Medida aplicable sin evidencias", f"{code} no referencia evidencias documentales.", "Enlazar procedimientos, registros o capturas que el auditor pueda verificar (CCN-STIC 808).", "CCN-STIC 802/808")
        if blank(r.get("responsable")):
            add("SOA-05", MENOR, code, "Medida aplicable sin responsable", f"{code} no tiene responsable asignado.", "Asignar un rol del art. 11 o del esquema de roles de la organización.", "RD 311/2022, art. 11–13")
        estado = str(r.get("estado") or "")
        p = pct_of(r)
        if re.search(r"parcial|planificad|en curso|pendiente", estado, re.I):
            txt = " ".join(str(r.get(k) or "") for k in ("observaciones", "evidencias", "refuerzos_elegidos", "mc_ref"))
            if not re.search(r"PTR|MC-\d+", txt, re.I):
                add("SOA-06", MENOR, code, "Implantación parcial sin plan de tratamiento", f"{code} figura como '{estado}' sin acción PTR ni medida compensatoria referenciada.",
                    "Registrar la acción pendiente en el Plan de Tratamiento de Riesgos con responsable y plazo.", "RD 311/2022, art. 14")
            if p >= 1:
                add("SOA-07", MENOR, code, "Porcentaje incoherente con el estado", f"{code} está '{estado}' pero declara {round(p * 100)} %.", "Alinear estado y porcentaje de implantación.", "Buenas prácticas SoA")
        elif re.fullmatch(r"implantada", estado, re.I) and p < 1:
            add("SOA-07", MENOR, code, "Porcentaje incoherente con el estado", f"{code} está 'Implantada' pero declara {round(p * 100)} %.", "Alinear estado y porcentaje de implantación.", "Buenas prácticas SoA")
        if re.search(r"compensada", str(r.get("aplica") or ""), re.I) and not any(code in CODE_RE.findall(str(m.get("medida") or "")) for m in mcs):
            add("MC-01", MAYOR, code, "Medida compensada sin registro de medida compensatoria", f"{code} se declara compensada pero ninguna MC del registro la sustituye.",
                "Registrar la medida compensatoria: ámbito, restricción, riesgo, medida, validación, mantenimiento y aprobación.", "RD 311/2022, art. 28.3")
        if pe["aplica"] and refuerzos is not None:
            tabla = [x for x in refuerzos if x["codigo"] == code]
            for rn in pe["obligatorios"]:
                row = next((x for x in tabla if x["refuerzo"] == rn), None)
                if not row or not re.match(r"^s[ií]", str(row["exigible"] or ""), re.I):
                    add("REF-01", MAYOR, f"{code} {rn}", "Refuerzo exigido no contemplado",
                        f"Con nivel {nivel}, {code} exige {rn}, pero la tabla de refuerzos {'lo marca como ' + repr(row['exigible']) if row else 'no lo recoge'}.",
                        "Incorporar el refuerzo a la implantación o justificar una medida compensatoria.", "RD 311/2022, Anexo II")
            for g in pe["grupos"]:
                if not any(x["refuerzo"] in g and re.search(r"alternativa|s[ií]", str(x["exigible"] or ""), re.I) and not blank(x["implementacion"]) for x in tabla):
                    add("REF-02", MENOR, f"{code} [{' o '.join(g)}]", "Grupo de refuerzos alternativos sin opción elegida",
                        f"{code} exige al menos uno de {', '.join(g)} y no consta cuál se ha implantado.", "Indicar la alternativa elegida y su implementación.", "RD 311/2022, Anexo II")
            requeridos = set(pe["obligatorios"]) | {x for g in pe["grupos"] for x in g}
            for x in tabla:
                if re.match(r"^s[ií]", str(x["exigible"] or ""), re.I) and x["refuerzo"] not in requeridos:
                    add("REF-03", OBS, f"{code} {x['refuerzo']}", "Refuerzo marcado exigible que ya no lo es",
                        f"Con nivel {nivel}, {code} no exige {x['refuerzo']}; la tabla lo sigue marcando como exigible.",
                        "Actualizar la tabla de refuerzos tras la (re)categorización; mantenerlo es sobrecumplimiento admisible.", "RD 311/2022, Anexo II")
    for code in by_code:
        if code not in ANEXO:
            add("STR-02", MENOR, code, "Código de medida desconocido", f"{code} no es una medida del Anexo II del RD 311/2022.", "Corregir el código (¿medida del RD 3/2010 derogado?).", "RD 311/2022, Anexo II")
    res.grado = sum(pcts) / len(pcts) if pcts else None

    # 3. Registro de medidas compensatorias
    for mc in mcs:
        faltan = [k for k in ("validacion", "mantenimiento", "aprobacion", "riesgo") if blank(mc.get(k))]
        if faltan:
            add("MC-02", MENOR, mc["id"], "Medida compensatoria incompleta", f"{mc['id']} no documenta: {', '.join(faltan)}.", "Completar la ficha de la medida compensatoria.", "RD 311/2022, art. 28.3")
        for c in CODE_RE.findall(str(mc.get("medida") or "")):
            r = by_code.get(c)
            if r and not re.search(r"compensada", f"{r.get('aplica') or ''} {r.get('estado') or ''}", re.I):
                add("MC-03", OBS, mc["id"], "Medida compensatoria no reflejada en la SoA", f"{mc['id']} sustituye a {c}, pero la SoA no la marca como compensada.",
                    "Marcar la medida como SÍ (compensada) y referenciar la MC.", "RD 311/2022, art. 28.3")

    # 4. Documento
    port = read_portada(ws_port)
    if port:
        firma = next((v for k, v in port.items() if k.lower().startswith("firma")), None)
        if blank(firma):
            add("DOC-01", MENOR, "Portada", "SoA sin firma", "La Declaración de Aplicabilidad no consta firmada por la Responsable de Seguridad.", "Firmar electrónicamente la SoA.", "RD 311/2022, art. 28.2")
        fe = next((v for k, v in port.items() if k.lower().startswith("fecha de emisi")), None)
        fecha = fe.date() if isinstance(fe, dt.datetime) else None
        if fecha is None and fe:
            m = re.search(r"(\d{2})/(\d{2})/(\d{4})", str(fe))
            fecha = dt.date(int(m[3]), int(m[2]), int(m[1])) if m else None
        if fecha and (hoy - fecha).days > 365:
            add("DOC-02", MENOR, "Portada", "SoA sin revisar en los últimos 12 meses", f"Fecha de emisión {fecha:%d/%m/%Y}.", "Revisar la SoA al menos anualmente.", "RD 311/2022, art. 28.2")

    res.hallazgos.sort(key=lambda f: (SEV_ORDER[f.sev], f.id, natural(f.ambito)))
    return res


def code_key(c):
    order = ["org", "op.pl", "op.acc", "op.ext", "op.nub", "op.exp", "op.cont", "op.mon", "mp.if", "mp.per", "mp.eq", "mp.com", "mp.si", "mp.sw", "mp.info", "mp.s"]
    fam = re.sub(r"\.\d+$", "", c)
    return (order.index(fam) if fam in order else 99, int(re.search(r"(\d+)$", c)[1]))


def natural(s):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", str(s))]


# ---------------------------------------------------------------- salida
def to_markdown(res: Result) -> str:
    L = [f"# Preauditoría de la SoA — {res.fichero}", "",
         f"**Fecha:** {res.fecha} · **Herramienta:** ens-soa-audit {__version__}", "",
         "## Resumen", "",
         f"- Categoría recalculada: **{res.categoria or 'no determinable'}** ({' · '.join(f'{d}={res.niveles.get(d) or chr(8212)}' for d in DIMS)}).",
         f"- Medidas en la SoA: {res.medidas} · exigidas por el Anexo II: {res.aplicables}" + (f" · grado de implantación declarado: {res.grado * 100:.1f} %." if res.grado is not None else "."),
         f"- Resultado: **{res.count(MAYOR)} NC mayores**, **{res.count(MENOR)} NC menores**, {res.count(OBS)} observaciones.", ""]
    for sev, title in ((MAYOR, "No conformidades mayores"), (MENOR, "No conformidades menores"), (OBS, "Observaciones")):
        fs = [f for f in res.hallazgos if f.sev == sev]
        if not fs:
            continue
        L += [f"## {title} ({len(fs)})", "", "| Regla | Ámbito | Hallazgo | Acción recomendada | Referencia |", "|---|---|---|---|---|"]
        L += [f"| {f.id} | {f.ambito} | **{f.titulo}.** {f.detalle.replace('|', '/')} | {f.recomendacion.replace('|', '/')} | {f.ref} |" for f in fs]
        L.append("")
    L.append("_Preauditoría automática: prepara la auditoría formal del art. 31 RD 311/2022, no la sustituye._")
    return "\n".join(L)


def print_console(res: Result, stream=sys.stdout):
    tty = stream.isatty()
    col = lambda s, c: f"\033[{c}m{s}\033[0m" if tty else s
    colors = {MAYOR: "31;1", MENOR: "33;1", OBS: "36"}
    print(col(f"ens-soa-audit {__version__} · {res.fichero}", "1"), file=stream)
    print(f"Categoría recalculada: {res.categoria} ({' '.join(f'{d}={res.niveles.get(d) or chr(8212)}' for d in DIMS)}) · "
          f"medidas {res.medidas} · exigidas {res.aplicables}" + (f" · implantación {res.grado * 100:.1f} %" if res.grado is not None else ""), file=stream)
    for f in res.hallazgos:
        print(f"  {col(f.sev.ljust(11), colors[f.sev])} {f.id:<7} {f.ambito:<18} {f.titulo}", file=stream)
        print(f"  {'':11} {'':7} {'':18} {f.detalle}", file=stream)
    print(col(f"\n{res.count(MAYOR)} NC mayores · {res.count(MENOR)} NC menores · {res.count(OBS)} observaciones", "1"), file=stream)


def main(argv=None):
    ap = argparse.ArgumentParser(prog="ens-soa-audit", description="Preauditoría de Declaraciones de Aplicabilidad del ENS (RD 311/2022).")
    ap.add_argument("xlsx", type=Path, help="Libro Excel con la SoA")
    ap.add_argument("--md", type=Path, help="Escribe el informe en Markdown")
    ap.add_argument("--json", type=Path, help="Escribe el resultado en JSON")
    ap.add_argument("--fecha", help="Fecha de referencia AAAA-MM-DD (por defecto, hoy)")
    ap.add_argument("--fail-on", choices=["mayor", "menor", "none"], default="none", help="Código de salida 2 si hay NC de esa severidad o superior")
    ap.add_argument("-q", "--quiet", action="store_true", help="No imprime el detalle en consola")
    ap.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    a = ap.parse_args(argv)
    if not a.xlsx.exists():
        ap.error(f"No existe {a.xlsx}")
    res = audit(a.xlsx, dt.date.fromisoformat(a.fecha) if a.fecha else None)
    if not a.quiet:
        print_console(res)
    if a.md:
        a.md.write_text(to_markdown(res), encoding="utf-8")
    if a.json:
        a.json.write_text(json.dumps({**asdict(res), "resumen": {MAYOR: res.count(MAYOR), MENOR: res.count(MENOR), OBS: res.count(OBS)}}, ensure_ascii=False, indent=1), encoding="utf-8")
    if a.fail_on == "mayor" and res.count(MAYOR):
        return 2
    if a.fail_on == "menor" and (res.count(MAYOR) or res.count(MENOR)):
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
